import io
import json
import logging
import os
import re
from datetime import datetime, timezone
from typing import Any, Dict, Optional, Tuple
import zipfile

import docx
import pypdf
from fastapi import HTTPException, status
from supabase import Client

logger = logging.getLogger(__name__)

RESUME_BUCKET = "resumes"
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MIN_TEXT_LENGTH = 20  # Minimum readable characters


def normalize_extracted_text(text: str) -> str:
    """Normalize extracted text by consolidating excessive whitespace while preserving structure."""
    if not text:
        return ""
    # Replace null bytes or non-printable chars
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    # Replace multiple spaces/tabs with single space
    text = re.sub(r"[ \t]+", " ", text)
    # Replace 3 or more consecutive newlines with 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def validate_and_extract_resume(
    filename: Optional[str],
    content_type: Optional[str],
    file_bytes: bytes,
) -> Tuple[str, str]:
    """
    Validate resume file format, size, content integrity, and extract text.
    Returns (cleaned_filename, extracted_text).
    Raises HTTPException on validation or extraction failure.
    """
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Resume file must have a valid filename.",
        )

    # 1. Enforce size limit
    if len(file_bytes) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Resume file exceeds maximum size limit of 5 MB.",
        )

    if len(file_bytes) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded resume file is empty.",
        )

    # 2. Enforce file extension
    ext = os.path.splitext(filename)[1].lower()
    if ext not in [".pdf", ".docx"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload a PDF (.pdf) or Word document (.docx).",
        )

    extracted_text = ""

    # 3. Content validation and extraction
    if ext == ".pdf":
        if not file_bytes.startswith(b"%PDF"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid PDF file. Header does not match standard PDF format.",
            )

        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if reader.is_encrypted:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Password-protected or encrypted PDF files are not supported.",
                )

            pages_text = []
            for page in reader.pages:
                page_txt = page.extract_text()
                if page_txt:
                    pages_text.append(page_txt)
            extracted_text = "\n\n".join(pages_text)
        except HTTPException:
            raise
        except Exception as exc:
            logger.warning("PDF extraction error: %s", exc)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Could not read or parse PDF file. The file may be corrupted.",
            )

    elif ext == ".docx":
        # Check zip magic bytes
        if not file_bytes.startswith(b"PK"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid Word document. File does not match DOCX packaging.",
            )

        try:
            # Verify zip contains word/document.xml
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                if "word/document.xml" not in zf.namelist():
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="Invalid DOCX package: missing Word document component.",
                    )

            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        paragraphs_text.append(row_txt)
            extracted_text = "\n\n".join(paragraphs_text)
        except HTTPException:
            raise
        except Exception as exc:
            logger.warning("DOCX extraction error: %s", exc)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Could not read or parse DOCX file. The file may be corrupted.",
            )

    normalized_text = normalize_extracted_text(extracted_text)

    # 4. Check for meaningful text content
    if len(normalized_text) < MIN_TEXT_LENGTH:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not extract readable text from document. Scanned image-only PDFs and empty documents are not supported.",
        )

    # Sanitize filename
    safe_name = os.path.basename(filename)
    safe_name = re.sub(r"[^a-zA-Z0-9._-]", "_", safe_name)

    return safe_name, normalized_text


def save_resume_and_metadata(
    client: Client,
    user_id: str,
    safe_filename: str,
    file_bytes: bytes,
    extracted_text: str,
) -> Dict[str, Any]:
    """
    Persist resume file and extracted text into Supabase Storage and database.
    Scoped strictly to the authenticated user_id.
    """
    ext = os.path.splitext(safe_filename)[1].lower()
    now_iso = datetime.now(timezone.utc).isoformat()
    resume_storage_path = f"{user_id}/resume{ext}"
    text_storage_path = f"{user_id}/resume_text.txt"
    meta_storage_path = f"{user_id}/resume_meta.json"

    content_type = "application/pdf" if ext == ".pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    # 1. Clean up existing files in user's storage folder to prevent orphaned files
    try:
        user_files = client.storage.from_(RESUME_BUCKET).list(user_id)
        if user_files:
            file_paths_to_remove = [f"{user_id}/{f['name']}" for f in user_files if isinstance(f, dict) and "name" in f]
            if file_paths_to_remove:
                client.storage.from_(RESUME_BUCKET).remove(file_paths_to_remove)
    except Exception as exc:
        logger.debug("Storage cleanup notice (non-fatal): %s", exc)

    # 2. Upload resume file
    try:
        client.storage.from_(RESUME_BUCKET).upload(
            resume_storage_path,
            file_bytes,
            {"content-type": content_type, "upsert": "true"},
        )
    except Exception as exc:
        logger.error("Failed to upload resume to Supabase Storage: %s", exc)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to store resume file.",
        )

    # 3. Store extracted text and metadata in private storage
    metadata_payload = {
        "user_id": user_id,
        "file_name": safe_filename,
        "file_size": len(file_bytes),
        "uploaded_at": now_iso,
        "storage_path": resume_storage_path,
    }

    try:
        client.storage.from_(RESUME_BUCKET).upload(
            text_storage_path,
            extracted_text.encode("utf-8"),
            {"content-type": "text/plain", "upsert": "true"},
        )
        client.storage.from_(RESUME_BUCKET).upload(
            meta_storage_path,
            json.dumps(metadata_payload).encode("utf-8"),
            {"content-type": "application/json", "upsert": "true"},
        )
    except Exception as exc:
        logger.warning("Could not persist text/meta to storage: %s", exc)

    # 4. Update profiles table
    try:
        # First try updating all resume columns
        client.table("profiles").update({
            "resume_available": True,
            "resume_file_path": resume_storage_path,
            "resume_file_name": safe_filename,
            "resume_uploaded_at": now_iso,
            "resume_text": extracted_text,
        }).eq("user_id", user_id).execute()
    except Exception:
        # Fallback if DB columns are not yet present: ensure resume_available is True
        try:
            client.table("profiles").update({
                "resume_available": True,
            }).eq("user_id", user_id).execute()
        except Exception as exc:
            logger.error("Failed to update profile resume_available: %s", exc)

    return metadata_payload


def get_user_resume_metadata(client: Client, user_id: str) -> Dict[str, Any]:
    """Retrieve metadata for the authenticated user's uploaded resume."""
    # Check profile
    profile_data = {}
    try:
        res = client.table("profiles").select("*").eq("user_id", user_id).execute()
        if res.data:
            profile_data = res.data[0]
    except Exception as exc:
        logger.error("Error querying profile for resume: %s", exc)

    if not profile_data.get("resume_available"):
        return {
            "resume_available": False,
            "file_name": None,
            "file_size": None,
            "uploaded_at": None,
        }

    # If DB profile has resume_file_name and resume_uploaded_at
    file_name = profile_data.get("resume_file_name")
    uploaded_at = profile_data.get("resume_uploaded_at")

    # If not in profile columns, read from storage resume_meta.json
    if not file_name or not uploaded_at:
        try:
            meta_bytes = client.storage.from_(RESUME_BUCKET).download(f"{user_id}/resume_meta.json")
            if meta_bytes:
                meta = json.loads(meta_bytes.decode("utf-8"))
                return {
                    "resume_available": True,
                    "file_name": meta.get("file_name"),
                    "file_size": meta.get("file_size"),
                    "uploaded_at": meta.get("uploaded_at"),
                }
        except Exception:
            pass

    return {
        "resume_available": True,
        "file_name": file_name or "resume.pdf",
        "file_size": None,
        "uploaded_at": uploaded_at or profile_data.get("updated_at"),
    }


def get_user_resume_text(client: Client, user_id: str) -> Optional[str]:
    """Retrieve extracted resume text for the authenticated user."""
    # 1. Try from profiles table
    try:
        res = client.table("profiles").select("resume_text, resume_available").eq("user_id", user_id).execute()
        if res.data:
            row = res.data[0]
            if not row.get("resume_available"):
                return None
            if row.get("resume_text"):
                return str(row["resume_text"])
    except Exception:
        pass

    # 2. Try from storage
    try:
        text_bytes = client.storage.from_(RESUME_BUCKET).download(f"{user_id}/resume_text.txt")
        if text_bytes:
            return text_bytes.decode("utf-8")
    except Exception:
        pass

    return None


def delete_user_resume(client: Client, user_id: str) -> bool:
    """Delete the authenticated user's resume from storage and clear profile metadata."""
    # 1. Remove storage files
    try:
        user_files = client.storage.from_(RESUME_BUCKET).list(user_id)
        if user_files:
            file_paths_to_remove = [f"{user_id}/{f['name']}" for f in user_files if isinstance(f, dict) and "name" in f]
            if file_paths_to_remove:
                client.storage.from_(RESUME_BUCKET).remove(file_paths_to_remove)
    except Exception as exc:
        logger.warning("Storage remove error: %s", exc)

    # 2. Update profiles table
    try:
        client.table("profiles").update({
            "resume_available": False,
            "resume_file_path": None,
            "resume_file_name": None,
            "resume_uploaded_at": None,
            "resume_text": None,
        }).eq("user_id", user_id).execute()
    except Exception:
        try:
            client.table("profiles").update({
                "resume_available": False,
            }).eq("user_id", user_id).execute()
        except Exception as exc:
            logger.error("Failed to update profile on resume delete: %s", exc)

    return True
